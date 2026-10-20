"""Unit tests for the report builder (report/build_report.py) on tiny Markdown samples.

The builder is a script, not part of the package, so it is loaded from its file. The tests cover
the caption convention ('Short title || explanation'), the table layout rules (repeated header
row, rows kept on one page, smaller text and landscape pages for wide tables) and the cached
contents and lists, without building the real report.
"""

from __future__ import annotations

import importlib.util
import sys
from pathlib import Path

import pytest
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn

BUILDER = Path(__file__).resolve().parents[1] / "report" / "build_report.py"


@pytest.fixture(scope="module")
def br():
    spec = importlib.util.spec_from_file_location("build_report", BUILDER)
    module = importlib.util.module_from_spec(spec)
    sys.modules["build_report"] = module  # dataclasses look the module up while the class is built
    spec.loader.exec_module(module)
    return module


def _document(br):
    document = Document()
    br.configure_styles(document)
    br.set_page_geometry(document.sections[0], landscape=False)
    return document


def _render(br, text: str, tmp_path: Path):
    """Number and render one Markdown sample the way build() does; return document and counters."""
    source = tmp_path / "01_sample.md"
    source.write_text(text, encoding="utf-8")
    counters = br.Counters()
    br.collect_numbers([source], counters)
    document = _document(br)
    br.render_markdown(document, text, counters, tmp_path)
    return document, counters


def _table(columns: int, rows: int = 2) -> str:
    lines = ["| " + " | ".join(f"H{i}" for i in range(columns)) + " |", "|" + "---|" * columns]
    lines += ["| " + " | ".join(f"r{r}c{i}" for i in range(columns)) + " |" for r in range(rows)]
    return "\n".join(lines)


def _orientations(document) -> list[str]:
    return ["L" if s.orientation == WD_ORIENT.LANDSCAPE else "P" for s in document.sections]


def test_split_caption_with_and_without_marker(br):
    assert br.split_caption("Short title || Longer explanation.") == ("Short title", "Longer explanation.")
    # the marker needs a space on both sides, so a bare '||' stays part of the title
    assert br.split_caption("Whole caption a||b") == ("Whole caption a||b", "")
    assert br.full_caption("Short title || Explanation.") == "Short title. Explanation."
    assert br.full_caption("Short title. || Explanation.") == "Short title. Explanation."
    assert br.full_caption("No marker here") == "No marker here"


def test_marker_splits_caption_into_short_title_and_small_explanation(br, tmp_path):
    text = (
        "As {tab:demo} shows, the sample works.\n\n"
        "Table: Short title || Longer explanation that cites {tab:demo} {#tab:demo}\n\n"
        + _table(3)
        + "\n"
    )
    document, counters = _render(br, text, tmp_path)
    paragraphs = document.paragraphs
    position = next(i for i, p in enumerate(paragraphs) if p.style.name == "Caption")
    caption, note = paragraphs[position], paragraphs[position + 1]
    assert caption.text == "Table 1. Short title"
    assert caption.paragraph_format.keep_with_next
    assert note.style.name == "Normal"
    assert note.text == "Longer explanation that cites Table 1"
    assert {run.font.size.pt for run in note.runs} == {br.EXPLANATION_PT}
    assert note.paragraph_format.keep_with_next  # the explanation stays with its table
    assert paragraphs[0].text == "As Table 1 shows, the sample works."
    assert counters.problems == []


def test_caption_without_marker_renders_as_before(br, tmp_path):
    text = "See {tab:demo}.\n\nTable: Whole caption, kept in one piece {#tab:demo}\n\n" + _table(3) + "\n"
    document, _ = _render(br, text, tmp_path)
    captions = [p for p in document.paragraphs if p.style.name == "Caption"]
    assert [p.text for p in captions] == ["Table 1. Whole caption, kept in one piece"]
    # the caption is directly followed by the table, with no explanation paragraph in between
    assert captions[0]._p.getnext().tag == qn("w:tbl")


def test_numbering_and_audit_accept_the_marker(br, tmp_path):
    source = tmp_path / "01_sample.md"
    source.write_text(
        "{fig:plot} and {tab:demo} are cited first.\n\n"
        "![Plot title || What the plot shows](missing.png){#fig:plot width=3}\n\n"
        "Table: Table title || What the table shows {#tab:demo}\n\n" + _table(2) + "\n",
        encoding="utf-8",
    )
    counters = br.Counters()
    br.collect_numbers([source], counters)
    assert counters.figures == {"plot": 1} and counters.tables == {"demo": 1}
    assert counters.figure_captions == ["Plot title || What the plot shows"]
    assert br.audit_sources([source], {}) == []


def test_every_table_repeats_its_header_and_keeps_rows_whole(br, tmp_path):
    text = "Table: Small {#tab:demo}\n\n" + _table(3, rows=3) + "\n\nSee {tab:demo}.\n"
    document, _ = _render(br, text, tmp_path)
    table = document.tables[0]
    rows = table._tbl.xpath("./w:tr")
    assert [bool(row.xpath("./w:trPr/w:tblHeader")) for row in rows] == [True, False, False, False]
    assert all(row.xpath("./w:trPr/w:cantSplit") for row in rows)
    width = table._tbl.tblPr.find(qn("w:tblW"))
    assert (width.get(qn("w:type")), width.get(qn("w:w"))) == ("pct", "5000")
    assert table._tbl.tblPr.find(qn("w:tblLayout")).get(qn("w:type")) == "autofit"
    assert {r.font.size.pt for c in table.rows[1].cells for p in c.paragraphs for r in p.runs if r.text} == {br.TABLE_PT}


@pytest.mark.parametrize(("columns", "size"), [(6, 11.0), (7, 9.0), (8, 8.5), (9, 8.5)])
def test_wide_tables_use_smaller_text(br, columns, size):
    assert br.table_font_size(columns) == size


def test_wide_table_gets_landscape_pages_and_returns_to_portrait(br, tmp_path):
    text = (
        "Before {tab:a}.\n\nTable: First wide {#tab:a}\n\n" + _table(9) + "\n\n"
        "Table: Second wide, sharing the pages {#tab:b}\n\n" + _table(10) + "\n\n"
        "After {tab:b}.\n\nTable: Narrow {#tab:c}\n\n" + _table(3) + "\n\nEnd {tab:c}.\n"
    )
    document, _ = _render(br, text, tmp_path)
    assert _orientations(document) == ["P", "L", "P"]
    sizes = [{r.font.size.pt for r in t.rows[1].cells[0].paragraphs[0].runs if r.text} for t in document.tables]
    assert sizes == [{8.5}, {8.5}, {br.TABLE_PT}]


def test_wide_table_inside_landscape_directive_adds_no_section(br, tmp_path):
    text = "See {tab:a}.\n\n\\landscape\n\nTable: Wide {#tab:a}\n\n" + _table(9) + "\n\n\\portrait\n\nAfter.\n"
    document, _ = _render(br, text, tmp_path)
    assert _orientations(document) == ["P", "L", "P"]


def test_page_break_after_wide_table_is_absorbed_by_the_section_break(br, tmp_path):
    text = "See {tab:a}.\n\nTable: Wide {#tab:a}\n\n" + _table(9) + "\n\n\\pagebreak\n\nAfter.\n"
    document, _ = _render(br, text, tmp_path)
    assert _orientations(document) == ["P", "L", "P"]
    assert not document.element.body.xpath(".//w:br[@w:type='page']")


def test_contents_and_lists_cache_real_entries(br, tmp_path):
    source = tmp_path / "01_sample.md"
    source.write_text("# 1 Introduction\n\nText.\n\n## 1.1 Business case\n\n### 1.1.1 Detail\n\n#### not listed\n",
                      encoding="utf-8")
    counters = br.Counters(
        figure_captions=["Plot title || What the plot shows"], table_captions=["Table title without marker"],
    )
    counters.headings = br.collect_headings([source])
    assert counters.headings == [(1, "1 Introduction"), (2, "1.1 Business case"), (3, "1.1.1 Detail")]
    document = _document(br)
    br.add_toc(document, counters)
    texts = [(p.style.name, p.text) for p in document.paragraphs]
    assert texts[:4] == [("Normal", "Table of contents"), ("toc 1", "1 Introduction"),
                         ("toc 2", "1.1 Business case"), ("toc 3", "1.1.1 Detail")]
    assert ("table of figures", "Figure 1. Plot title") in texts
    assert ("table of figures", "Table 1. Table title without marker") in texts
    assert not any("Update Field" in text for _, text in texts)
    body = document.element.body
    instructions = [t.text.strip() for t in body.iter(qn("w:instrText"))]
    assert instructions == ['TOC \\o "1-3" \\h \\z \\u', 'TOC \\h \\z \\c "Figure"', 'TOC \\h \\z \\c "Table"']
    # the contents field begins in its first entry and ends in its last
    first, last = document.paragraphs[1]._p, document.paragraphs[3]._p
    assert first.xpath("./w:r[1]/w:fldChar/@w:fldCharType") == ["begin", "separate"]
    assert last.xpath("./w:r[last()]/w:fldChar/@w:fldCharType") == ["end"]
    assert len(body.xpath(".//w:fldChar[@w:fldCharType='begin']")) == len(body.xpath(".//w:fldChar[@w:fldCharType='end']"))


def test_update_fields_setting_sits_in_schema_order(br):
    document = Document()
    br.set_update_fields_on_open(document)
    br.set_update_fields_on_open(document)  # idempotent
    settings = document.settings.element
    nodes = settings.findall(qn("w:updateFields"))
    assert len(nodes) == 1 and nodes[0].get(qn("w:val")) == "true"
    assert nodes[0].getnext().tag == qn("w:compat")
