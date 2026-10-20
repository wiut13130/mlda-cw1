"""Build the coursework report (.docx) from Markdown section files.

Usage (from the repository root):
    uv run python report/build_report.py            # writes report/MLDA_CW1_13130_report.docx
    uv run python report/build_report.py --check    # prints word counts and source problems only (exit 1 if any)

Why a generator rather than a hand-edited document: every number in the report is produced by the
pipeline, so the text is kept in version control as Markdown (report/sections/*.md) and rendered
with the formatting rules of the brief (Times New Roman 12, single spacing, A4, student ID in the
header of every page, page numbers, automatic table of contents, Harvard references). Re-running
the pipeline and this script regenerates a consistent document.

Supported Markdown subset:
    # / ## / ### headings            paragraphs            - bullets   1. numbered items
    **bold** and *italic* inline    pipe tables (| a | b |) preceded by an optional
    'Table: caption text {#tab:key}' line
    ![Caption text](reports/figures/file.png){#fig:key width=6.3}   (the caption is also the alt text)
    'Short title || explanation' in a figure or table caption: the short title goes into the
    Caption paragraph (and so into the lists of figures and tables), the explanation into a
    smaller paragraph under it; a caption without the marker is all title
    {fig:key} / {tab:key} cross-references replaced by the assigned numbers
    \\pagebreak                        > quote (rendered as an indented paragraph)
    \\landscape / \\portrait          start a landscape or a portrait section (for a wide table)
Every table repeats its header row on each page and keeps each row on one page; tables of seven or
more columns use smaller text, and a table of nine or more columns in a portrait section is put on
landscape pages of its own automatically.
Front matter (cover page, table of contents, lists of figures and tables, abbreviations) is generated
from report/report_meta.json; sections render in file-name order; files whose name starts with
"9" are treated as references (90_references.md) or appendices (9A_, 9B_, ...).
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
REPORT_DIR = ROOT / "report"
SECTIONS_DIR = REPORT_DIR / "sections"
META_PATH = REPORT_DIR / "report_meta.json"
FONT = "Times New Roman"
BODY_PT = 12
CAPTION_PT = 12
TABLE_PT = 11
EXPLANATION_PT = 9  # caption explanation under the short title
MARGIN_CM = 2.54
TEXT_WIDTH_IN = 6.27  # A4 width 21.0 cm minus 2 x 2.54 cm margins
# Column counts from which a table gets smaller text and, in a portrait section, landscape pages:
# at 11 pt seven or more columns of a 16 cm text width wrap numbers and headers onto several lines.
SMALL_TABLE_COLUMNS = 7
SMALLER_TABLE_COLUMNS = 8
LANDSCAPE_TABLE_COLUMNS = 9

FIG_RE = re.compile(r"^!\[(?P<caption>.*?)\]\((?P<path>[^)]+)\)(?:\{(?P<attrs>[^}]*)\})?\s*$")
TAB_CAPTION_RE = re.compile(r"^Table:\s*(?P<caption>.*?)(?:\s*\{#tab:(?P<key>[\w-]+)\})?\s*$")
HEADING_RE = re.compile(r"^(#{1,3})\s+(.*)$")
XREF_RE = re.compile(r"\{(fig|tab):([\w-]+)\}")
INLINE_RE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)")
PLACEHOLDER_RE = re.compile(r"\{\{[^{}]*\}\}|\[\[[A-Z0-9_]+\]\]")   # {{name}} and [[NAME]] tokens block a build
CAPTION_MARKER = " || "

# Children of w:settings that the schema (CT_Settings is a sequence) places after w:updateFields;
# the element must go before the first of them that is present, not simply at the end.
SETTINGS_AFTER_UPDATE_FIELDS = (
    "w:hdrShapeDefaults", "w:footnotePr", "w:endnotePr", "w:compat", "w:docVars", "w:rsids",
    "m:mathPr", "w:attachedSchema", "w:themeFontLang", "w:clrSchemeMapping",
    "w:doNotIncludeSubdocsInStats", "w:doNotAutoCompressPictures", "w:forceUpgrade", "w:captions",
    "w:readModeInkLockDown", "w:smartTagType", "sl:schemaLibrary", "w:shapeDefaults",
    "w:doNotEmbedSmartTags", "w:decimalSymbol", "w:listSeparator",
)
# Children of w:tblPr that the schema places after w:tblW.
TBLPR_AFTER_TBLW = (
    "w:jc", "w:tblCellSpacing", "w:tblInd", "w:tblBorders", "w:shd", "w:tblLayout", "w:tblCellMar",
    "w:tblLook", "w:tblCaption", "w:tblDescription", "w:tblPrChange",
)


@dataclass
class Counters:
    figures: dict[str, int] = field(default_factory=dict)
    tables: dict[str, int] = field(default_factory=dict)
    figure_captions: list[str] = field(default_factory=list)
    table_captions: list[str] = field(default_factory=list)
    # (level, text) of every level-1 to level-3 heading in rendering order, for the cached contents.
    headings: list[tuple[int, str]] = field(default_factory=list)
    problems: list[str] = field(default_factory=list)
    # True right after a page or section break, so that a level-1 heading does not add a second
    # break and leave a blank page (the front matter ends with a break before the first section).
    fresh_page: bool = False
    # True while the renderer is in a landscape section it opened itself for a wide table; the
    # return to portrait waits for the next block that is not another wide table, so consecutive
    # wide tables share their landscape pages instead of leaving an empty portrait page between.
    auto_landscape: bool = False


# ----------------------------------------------------------------------------- helpers: OOXML


def _field_code(instruction: str) -> list:
    """The begin mark (flagged dirty so Word refreshes it), the instruction and the separate mark."""
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin.set(qn("w:dirty"), "true")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" {instruction} "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    return [begin, instr, separate]


def _field_end():
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    return end


def add_field(paragraph, instruction: str, cached: str = "") -> None:
    """Insert a complex field (PAGE, NUMPAGES, SEQ, TOC) with a cached result Word can refresh."""
    run = paragraph.add_run()
    run._r.extend(_field_code(instruction))
    result = paragraph.add_run(cached)
    end_run = paragraph.add_run()
    end_run._r.append(_field_end())
    for r in (run, result, end_run):
        r.font.name = FONT
        r.font.size = Pt(BODY_PT)


def wrap_in_field(paragraphs: list, instruction: str) -> None:
    """Make already-filled paragraphs the cached result of one complex field (a TOC field).

    Viewers that do not update fields (browser previews, most mobile apps, PDF converters) show
    the cached result, so caching the real entries gives a usable list everywhere, while Word can
    still rebuild it from the instruction with page numbers. The field begins in the first
    paragraph and ends in the last, which is how Word itself stores a multi-paragraph TOC.
    """
    first, last = paragraphs[0], paragraphs[-1]
    begin_run = first.add_run()
    begin_run._r.extend(_field_code(instruction))
    properties = first._p.pPr
    if properties is None:
        first._p.insert(0, begin_run._r)  # lxml moves the run rather than copying it
    else:
        properties.addnext(begin_run._r)
    end_run = last.add_run()
    end_run._r.append(_field_end())
    for run in (begin_run, end_run):
        run.font.name = FONT
        run.font.size = Pt(BODY_PT)


def set_update_fields_on_open(document) -> None:
    """Ask Word to refresh every field (contents, lists, SEQ numbers, page counts) on opening."""
    settings = document.settings.element
    node = settings.find(qn("w:updateFields"))
    if node is None:
        node = OxmlElement("w:updateFields")
        settings.insert_element_before(node, *SETTINGS_AFTER_UPDATE_FIELDS)
    node.set(qn("w:val"), "true")


def set_cell_shading(cell, hex_fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), hex_fill)
    properties.append(shading)


def keep_with_next(paragraph) -> None:
    paragraph.paragraph_format.keep_with_next = True


def set_row_flags(table, header_rows: int = 1) -> None:
    """Repeat the first ``header_rows`` rows on every page and keep each row on one page.

    A row split across a page break leaves half a model's scores on each page, and a long table
    continued without its header row cannot be read on its own.
    """
    for index, row in enumerate(table.rows):
        properties = row._tr.get_or_add_trPr()
        properties.append(OxmlElement("w:cantSplit"))
        if index < header_rows:
            properties.append(OxmlElement("w:tblHeader"))


def fit_table_to_window(table) -> None:
    """Word's AutoFit to Window: the table spans the text width and columns follow their content."""
    table.autofit = True
    properties = table._tbl.tblPr
    width = properties.find(qn("w:tblW"))
    if width is None:
        width = OxmlElement("w:tblW")
        properties.insert_element_before(width, *TBLPR_AFTER_TBLW)
    width.set(qn("w:type"), "pct")
    width.set(qn("w:w"), "5000")  # fiftieths of a percent, i.e. 100% of the text width


def table_font_size(columns: int) -> float:
    """Smaller text for wide tables, so that numbers and headers do not wrap one word per line."""
    if columns >= SMALLER_TABLE_COLUMNS:
        return 8.5
    if columns >= SMALL_TABLE_COLUMNS:
        return 9
    return TABLE_PT


# ----------------------------------------------------------------------------- styles and layout


def configure_styles(document) -> None:
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(BODY_PT)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.line_spacing = 1.0
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for level, size in ((1, 16), (2, 14), (3, 12)):
        style = styles[f"Heading {level}"]
        style.font.name = FONT
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.paragraph_format.space_before = Pt(18 if level == 1 else 12)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    caption = styles["Caption"]
    caption.font.name = FONT
    caption.font.size = Pt(CAPTION_PT)
    caption.font.bold = False
    caption.font.italic = False
    caption.font.color.rgb = RGBColor(0, 0, 0)
    caption.paragraph_format.space_after = Pt(10)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for name in ("List Bullet", "List Number"):
        style = styles[name]
        style.font.name = FONT
        style.font.size = Pt(BODY_PT)
    # Word's built-in styles for contents entries and list-of-figures entries. The cached entries
    # use them and Word applies them when it refreshes the fields, so the front matter looks the
    # same before and after an update. The template only has them as latent styles, so they are
    # defined here (as built-in, by their Word names).
    for name, indent_cm, space_before in (("toc 1", 0.0, 4), ("toc 2", 0.6, 0), ("toc 3", 1.2, 0)):
        style = _builtin_paragraph_style(styles, name)
        style.paragraph_format.left_indent = Cm(indent_cm)
        style.paragraph_format.space_before = Pt(space_before)
        style.paragraph_format.space_after = Pt(2)
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    figures = _builtin_paragraph_style(styles, "table of figures")
    figures.paragraph_format.left_indent = Cm(0.8)
    figures.paragraph_format.first_line_indent = Cm(-0.8)
    figures.paragraph_format.space_after = Pt(2)


def _builtin_paragraph_style(styles, name: str):
    if name in styles:
        return styles[name]
    style = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH, builtin=True)
    style.base_style = styles["Normal"]
    style.next_paragraph_style = styles["Normal"]
    style.unhide_when_used = True
    return style


def set_page_geometry(section, landscape: bool = False) -> None:
    """A4 in portrait or landscape orientation with 2.54 cm margins."""
    section.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    long_side, short_side = Cm(29.7), Cm(21.0)
    section.page_width, section.page_height = (long_side, short_side) if landscape else (short_side, long_side)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Cm(MARGIN_CM))


def start_section(document, landscape: bool, counters: Counters) -> None:
    """Start a new page in a new section; its header and footer are inherited from the first section."""
    section = document.add_section(WD_SECTION.NEW_PAGE)
    set_page_geometry(section, landscape)
    counters.fresh_page = True


def page_break(document, counters: Counters) -> None:
    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    counters.fresh_page = True


def configure_page(document, student_id: str) -> None:
    section = document.sections[0]
    set_page_geometry(section, landscape=False)
    section.different_first_page_header_footer = False
    header = section.header.paragraphs[0]
    header.text = ""
    run = header.add_run(f"Student ID: {student_id}    6COSC017C Machine Learning and Data Analytics, CW1")
    run.font.name = FONT
    run.font.size = Pt(10)
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer = section.footer.paragraphs[0]
    footer.text = ""
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    lead = footer.add_run("Page ")
    lead.font.name = FONT
    lead.font.size = Pt(10)
    add_field(footer, "PAGE", "1")
    mid = footer.add_run(" of ")
    mid.font.name = FONT
    mid.font.size = Pt(10)
    add_field(footer, "NUMPAGES", "1")


# ----------------------------------------------------------------------------- inline markdown


def add_inline(paragraph, text: str, counters: Counters, size: float | None = None) -> None:
    """Add text with **bold**, *italic*, `code` and {fig:key}/{tab:key} substitution."""
    text = XREF_RE.sub(lambda m: _resolve_xref(m, counters), text)
    for token in INLINE_RE.split(text):
        if not token:
            continue
        if token.startswith("**") and token.endswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("*") and token.endswith("*"):
            run = paragraph.add_run(token[1:-1])
            run.italic = True
        elif token.startswith("`") and token.endswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
        else:
            run = paragraph.add_run(token)
        if run.font.name is None:
            run.font.name = FONT
        run.font.size = Pt(size or BODY_PT)


def _resolve_xref(match: re.Match, counters: Counters) -> str:
    kind, key = match.group(1), match.group(2)
    table = counters.figures if kind == "fig" else counters.tables
    if key not in table:
        counters.problems.append(f"unresolved cross-reference {{{kind}:{key}}}")
        return f"{'Figure' if kind == 'fig' else 'Table'} ??"
    return f"{'Figure' if kind == 'fig' else 'Table'} {table[key]}"


def plain_text(text: str, counters: Counters) -> str:
    """Caption text with cross-references resolved and Markdown emphasis removed (for alt text)."""
    resolved = XREF_RE.sub(lambda m: _resolve_xref(m, counters), text)
    return re.sub(r"(\*\*|\*|`)", "", resolved)


def split_caption(caption: str) -> tuple[str, str]:
    """Split 'Short title || explanation' into its two parts.

    A caption without the marker is all title, so captions written before the convention keep
    rendering exactly as they did.
    """
    title, marker, explanation = caption.partition(CAPTION_MARKER)
    if not marker:
        return caption, ""
    return title.strip(), explanation.strip()


def full_caption(caption: str) -> str:
    """Title and explanation run together as sentences, for the alternative text of a figure."""
    title, explanation = split_caption(caption)
    if not explanation:
        return title
    joiner = " " if title.endswith((".", "!", "?", ":", ";")) else ". "
    return f"{title}{joiner}{explanation}"


def set_alt_text(picture, description: str) -> None:
    """Write the alternative text of an inline picture (Word: Format Picture, Alt Text)."""
    doc_pr = picture._inline.docPr
    doc_pr.set("descr", description)
    doc_pr.set("title", description.split(". ", 1)[0])


# ----------------------------------------------------------------------------- first pass: numbering


def collect_numbers(files: list[Path], counters: Counters) -> None:
    """Assign figure and table numbers in document order so cross-references can be resolved."""
    for path in files:
        lines = path.read_text(encoding="utf-8").splitlines()
        for line_no, line in enumerate(lines, start=1):
            fig = FIG_RE.match(line.strip())
            if fig:
                attrs = fig.group("attrs") or ""
                key_match = re.search(r"#fig:([\w-]+)", attrs)
                key = key_match.group(1) if key_match else f"auto{len(counters.figure_captions) + 1}"
                if key in counters.figures:
                    counters.problems.append(f"duplicate figure key fig:{key} ({path.name}:{line_no})")
                counters.figures[key] = len(counters.figure_captions) + 1
                counters.figure_captions.append(fig.group("caption"))
                continue
            tab = TAB_CAPTION_RE.match(line.strip())
            if tab:
                key = tab.group("key") or f"auto{len(counters.table_captions) + 1}"
                if key in counters.tables:
                    counters.problems.append(f"duplicate table key tab:{key} ({path.name}:{line_no})")
                counters.tables[key] = len(counters.table_captions) + 1
                counters.table_captions.append(tab.group("caption"))


def collect_headings(files: list[Path]) -> list[tuple[int, str]]:
    """Level and text of every level-1 to level-3 heading, in rendering order.

    The contents page is written before the sections are rendered, so its cached entries come
    from this pass; the pattern is the renderer's own, so both see the same headings.
    """
    headings: list[tuple[int, str]] = []
    for path in files:
        for line in path.read_text(encoding="utf-8").splitlines():
            match = HEADING_RE.match(line.strip())
            if match:
                headings.append((len(match.group(1)), match.group(2)))
    return headings


def audit_sources(files: list[Path], meta: dict) -> list[str]:
    """Problems that rendering alone does not reveal, for the files in rendering order.

    A figure or table must be referenced before it appears and at least once; ``{{...}}``
    placeholders must be filled; cover-sheet values in report_meta.json must not be stubs.
    """
    problems: list[str] = []
    defined: dict[tuple[str, str], tuple[int, int, str]] = {}
    first_reference: dict[tuple[str, str], tuple[int, int, str]] = {}
    for file_index, path in enumerate(files):
        for line_no, line in enumerate(path.read_text(encoding="utf-8").splitlines(), start=1):
            position = (file_index, line_no, path.name)
            for match in XREF_RE.finditer(line):
                first_reference.setdefault((match.group(1), match.group(2)), position)
            stripped = line.strip()
            fig = FIG_RE.match(stripped)
            key_match = re.search(r"#fig:([\w-]+)", fig.group("attrs") or "") if fig else None
            if key_match:
                defined.setdefault(("fig", key_match.group(1)), position)
            tab = TAB_CAPTION_RE.match(stripped)
            if tab and tab.group("key"):
                defined.setdefault(("tab", tab.group("key")), position)
            for match in PLACEHOLDER_RE.finditer(line):
                problems.append(f"unfilled placeholder {match.group(0)} ({path.name}:{line_no})")
    for (kind, key), (file_index, line_no, name) in defined.items():
        reference = first_reference.get((kind, key))
        if reference is None:
            problems.append(f"{kind}:{key} is never referenced ({name}:{line_no})")
        elif reference[:2] > (file_index, line_no):
            problems.append(
                f"{kind}:{key} is first referenced at {reference[2]}:{reference[1]}, after it appears at {name}:{line_no}"
            )
    for key, value in meta.items():
        if isinstance(value, str) and (PLACEHOLDER_RE.search(value) or value.lower().startswith("to be added")):
            problems.append(f"report_meta.json '{key}' is not filled: {value!r}")
    return problems


# ----------------------------------------------------------------------------- block rendering


def add_caption(document, label: str, number: int, text: str, counters: Counters):
    """Write 'Label n. Short title' (Caption style, SEQ field) and, if any, the explanation.

    Word builds its lists of figures and tables from Caption paragraphs, so the explanation after
    the ' || ' marker goes into a separate, smaller Normal paragraph and the lists show only the
    short titles. Returns the last paragraph written, which a table caption keeps with its table.
    """
    title, explanation = split_caption(text)
    paragraph = document.add_paragraph(style="Caption")
    lead = paragraph.add_run(f"{label} ")
    lead.font.name = FONT
    lead.font.size = Pt(CAPTION_PT)
    add_field(paragraph, f"SEQ {label} \\* ARABIC", str(number))
    tail = paragraph.add_run(". ")
    tail.font.name = FONT
    tail.font.size = Pt(CAPTION_PT)
    add_inline(paragraph, title, counters, size=CAPTION_PT)
    if not explanation:
        return paragraph
    paragraph.paragraph_format.space_after = Pt(2)
    keep_with_next(paragraph)
    note = document.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.LEFT
    note.paragraph_format.space_after = Pt(10)  # the gap the Caption style leaves after a caption
    add_inline(note, explanation, counters, size=EXPLANATION_PT)
    return note


def render_table(document, header: list[str], rows: list[list[str]], counters: Counters) -> None:
    size = table_font_size(len(header))
    table = document.add_table(rows=1, cols=len(header))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for index, title in enumerate(header):
        cell = table.rows[0].cells[index]
        cell.text = ""
        add_inline(cell.paragraphs[0], title.strip(), counters, size=size)
        for run in cell.paragraphs[0].runs:
            run.bold = True
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
        set_cell_shading(cell, "E7E6E6")
    for row in rows:
        cells = table.add_row().cells
        for index in range(len(header)):
            value = row[index].strip() if index < len(row) else ""
            cells[index].text = ""
            add_inline(cells[index].paragraphs[0], value, counters, size=size)
            cells[index].paragraphs[0].paragraph_format.space_after = Pt(0)
            cells[index].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_row_flags(table, header_rows=1)
    fit_table_to_window(table)
    document.add_paragraph().paragraph_format.space_after = Pt(4)


def split_table_row(line: str) -> list[str]:
    return [c for c in line.strip().strip("|").split("|")]


def parse_table_block(lines: list[str]) -> tuple[list[str], list[list[str]]]:
    header = split_table_row(lines[0])
    body = [split_table_row(line) for line in lines[2:] if line.strip()]
    return header, body


def is_wide_table_line(line: str) -> bool:
    """A table row with enough columns to need landscape pages."""
    return line.startswith("|") and len(split_table_row(line)) >= LANDSCAPE_TABLE_COLUMNS


def in_portrait_section(document) -> bool:
    """Orientation of the section being written (the last one), read from the document itself."""
    return document.sections[-1].orientation == WD_ORIENT.PORTRAIT


def end_auto_landscape(document, line: str, counters: Counters) -> bool:
    """Leave the landscape section opened for wide tables before rendering ``line``.

    Returns True when ``line`` (a page break) is used up by the section break, which already
    starts a new page; a \\landscape or \\portrait directive starts its own section instead.
    """
    counters.auto_landscape = False
    if line in ("\\landscape", "\\portrait"):
        return False
    start_section(document, False, counters)
    return line == "\\pagebreak"


def render_markdown(document, text: str, counters: Counters, figure_root: Path) -> int:
    """Render one Markdown file; return the number of body words (headings included, captions excluded)."""
    lines = text.splitlines()
    words = 0
    index = 0
    pending_table_caption: tuple[str, str | None] | None = None
    while index < len(lines):
        line = lines[index]
        stripped = line.strip()
        if not stripped:
            index += 1
            continue
        # A table caption only sets the pending caption, and another wide table stays on the
        # landscape pages; anything else ends the landscape section opened for wide tables.
        if counters.auto_landscape and not (TAB_CAPTION_RE.match(stripped) or is_wide_table_line(stripped)):
            if end_auto_landscape(document, stripped, counters):
                index += 1
                continue
        if stripped == "\\pagebreak":
            page_break(document, counters)
            index += 1
            continue
        if stripped in ("\\landscape", "\\portrait"):
            start_section(document, stripped == "\\landscape", counters)
            index += 1
            continue
        heading = HEADING_RE.match(stripped)
        if heading:
            level = len(heading.group(1))
            if level == 1 and not counters.fresh_page:
                page_break(document, counters)
            counters.fresh_page = False
            paragraph = document.add_heading(level=level)
            add_inline(paragraph, heading.group(2), counters, size={1: 16, 2: 14, 3: 12}[level])
            for run in paragraph.runs:
                run.bold = True
                run.font.color.rgb = RGBColor(0, 0, 0)
            words += len(heading.group(2).split())
            index += 1
            continue
        fig = FIG_RE.match(stripped)
        if fig:
            attrs = fig.group("attrs") or ""
            key_match = re.search(r"#fig:([\w-]+)", attrs)
            width_match = re.search(r"width=([\d.]+)", attrs)
            width = min(float(width_match.group(1)), TEXT_WIDTH_IN) if width_match else TEXT_WIDTH_IN
            key = key_match.group(1) if key_match else None
            number = counters.figures.get(key) if key else None
            if number is None:
                number = counters.figure_captions.index(fig.group("caption")) + 1
            image_path = (figure_root / fig.group("path")).resolve()
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            keep_with_next(paragraph)
            counters.fresh_page = False
            if image_path.exists():
                picture = paragraph.add_run().add_picture(str(image_path), width=Cm(width * 2.54))
                set_alt_text(picture, f"Figure {number}. {plain_text(full_caption(fig.group('caption')), counters)}")
            else:
                counters.problems.append(f"missing figure file {image_path}")
                paragraph.add_run(f"[missing figure: {fig.group('path')}]")
            add_caption(document, "Figure", number, fig.group("caption"), counters)
            index += 1
            continue
        tab = TAB_CAPTION_RE.match(stripped)
        if tab:
            pending_table_caption = (tab.group("caption"), tab.group("key"))
            index += 1
            continue
        counters.fresh_page = False
        if stripped.startswith("|"):
            block = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                block.append(lines[index])
                index += 1
            header, rows = parse_table_block(block)
            if len(header) >= LANDSCAPE_TABLE_COLUMNS and in_portrait_section(document):
                # Landscape pages of its own, unless a \landscape directive is already in force
                # (then the section is not portrait) or an earlier wide table opened them.
                start_section(document, True, counters)
                counters.auto_landscape = True
                counters.fresh_page = False
            if pending_table_caption:
                caption, key = pending_table_caption
                number = counters.tables.get(key) if key else counters.table_captions.index(caption) + 1
                caption_paragraph_added = True
                keep_with_next(add_caption(document, "Table", number, caption, counters))
                pending_table_caption = None
            else:
                caption_paragraph_added = False
                counters.problems.append(f"table without caption near: {block[0][:60]}")
            render_table(document, header, rows, counters)
            _ = caption_paragraph_added
            continue
        bullet = re.match(r"^[-*]\s+(.*)$", stripped)
        numbered = re.match(r"^\d+[.)]\s+(.*)$", stripped)
        if bullet or numbered:
            style = "List Bullet" if bullet else "List Number"
            content = (bullet or numbered).group(1)
            paragraph = document.add_paragraph(style=style)
            paragraph.paragraph_format.space_after = Pt(3)
            add_inline(paragraph, content, counters)
            words += len(content.split())
            index += 1
            continue
        if stripped.startswith(">"):
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Cm(1.0)
            add_inline(paragraph, stripped.lstrip("> "), counters)
            words += len(stripped.split())
            index += 1
            continue
        # paragraph: join consecutive non-empty, non-special lines
        block = [stripped]
        index += 1
        while index < len(lines):
            nxt = lines[index].strip()
            if (not nxt or nxt.startswith(("#", "|", "![", "- ", "* ", ">", "Table:", "\\pagebreak", "\\landscape", "\\portrait"))
                    or re.match(r"^\d+[.)]\s", nxt)):
                break
            block.append(nxt)
            index += 1
        paragraph = document.add_paragraph()
        content = " ".join(block)
        add_inline(paragraph, content, counters)
        words += len(XREF_RE.sub("Figure 0", content).split())
    return words


# ----------------------------------------------------------------------------- front matter


def add_cover(document, meta: dict, word_count: int) -> None:
    def centred(text: str, size: int, bold: bool = False, space_after: int = 12) -> None:
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(space_after)
        run = paragraph.add_run(text)
        run.font.name = FONT
        run.font.size = Pt(size)
        run.bold = bold

    for _ in range(5):
        document.add_paragraph()
    centred(meta["university"], 14, space_after=6)
    centred(meta["school"], 12, space_after=36)
    centred(f"{meta['module_code']} {meta['module_title']}", 14, bold=True, space_after=6)
    centred(meta["assessment"], 13, space_after=36)
    centred(meta["title"], 18, bold=True, space_after=6)
    centred(meta["subtitle"], 13, space_after=48)
    for label, key in (
        ("Student ID", "student_id"),
        ("Marker", "marker"),
        ("Submission date", "submission_date"),
        ("Word count (executive summary and sections 1 to 5 with headings; tables, captions, references "
         "and appendices excluded)", None),
        ("Repository", "repo_url"),
        ("Deployed application", "app_url"),
    ):
        value = f"{word_count:,}" if key is None else meta.get(key, "")
        centred(f"{label}: {value}", 12, space_after=4)
    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def add_toc(document, counters: Counters) -> None:
    """Contents, list of figures and list of tables, each a TOC field with a real cached list.

    The cached entries (headings with their numbers as written, and 'Figure n. Short title') are
    what any viewer shows; Word replaces them with the same entries plus page numbers when it
    updates the fields, which the document settings ask it to do on opening.
    """
    heading = document.add_paragraph()
    run = heading.add_run("Table of contents")
    run.bold = True
    run.font.size = Pt(16)
    run.font.name = FONT
    entries = []
    for level, text in counters.headings:
        entry = document.add_paragraph(style=f"toc {level}")
        add_inline(entry, text, counters)
        entries.append(entry)
    wrap_in_field(entries or [document.add_paragraph()], 'TOC \\o "1-3" \\h \\z \\u')
    document.add_paragraph()
    for label, kind, captions in (
        ("List of figures", "Figure", counters.figure_captions),
        ("List of tables", "Table", counters.table_captions),
    ):
        head = document.add_paragraph()
        run = head.add_run(label)
        run.bold = True
        run.font.size = Pt(14)
        run.font.name = FONT
        items = []
        for number, caption in enumerate(captions, start=1):
            item = document.add_paragraph(style="table of figures")
            add_inline(item, f"{kind} {number}. {split_caption(caption)[0]}", counters, size=CAPTION_PT)
            items.append(item)
        if items:
            wrap_in_field(items, f'TOC \\h \\z \\c "{kind}"')
        document.add_paragraph()
    page_break(document, counters)


def add_abbreviations(document, meta: dict, counters: Counters) -> None:
    items = meta.get("abbreviations", {})
    if not items:
        return
    heading = document.add_paragraph()
    run = heading.add_run("Abbreviations")
    run.bold = True
    run.font.size = Pt(14)
    run.font.name = FONT
    table = document.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    for abbreviation, meaning in items.items():
        cells = table.add_row().cells
        for cell, text in zip(cells, (abbreviation, meaning), strict=True):
            cell.text = ""
            add_inline(cell.paragraphs[0], text, counters, size=TABLE_PT)
            cell.paragraphs[0].paragraph_format.space_after = Pt(0)
    set_row_flags(table, header_rows=0)  # no header row: every row is an abbreviation
    fit_table_to_window(table)
    page_break(document, counters)


# ----------------------------------------------------------------------------- main


def section_files() -> tuple[list[Path], list[Path], list[Path]]:
    files = sorted(SECTIONS_DIR.glob("*.md"))
    body = [f for f in files if not f.name.startswith("9")]
    references = [f for f in files if f.name.startswith("90")]
    appendices = [f for f in files if f.name.startswith("9") and f not in references]
    return body, references, appendices


def build(output: Path | None = None, check_only: bool = False) -> Path:
    meta = json.loads(META_PATH.read_text(encoding="utf-8"))
    body, references, appendices = section_files()
    if not body:
        raise SystemExit(f"no section files found in {SECTIONS_DIR}")
    counters = Counters()
    collect_numbers(body + appendices, counters)
    counters.headings = collect_headings(body + references + appendices)
    source_problems = counters.problems + audit_sources(body + references + appendices, meta)
    counters.problems = []

    # word count pass on a scratch document (cheap; the renderer is the single source of truth)
    scratch = Document()
    configure_styles(scratch)
    body_words = sum(render_markdown(scratch, f.read_text(encoding="utf-8"), counters, ROOT) for f in body)
    counters.problems.clear()
    counters.auto_landscape = False  # the scratch pass may end on a wide table; the report starts portrait

    if check_only:
        print(f"body words: {body_words:,}")
        problems = list(source_problems)
        for path in body + references + appendices:
            file_counters = Counters(
                figures=counters.figures, tables=counters.tables,
                figure_captions=counters.figure_captions, table_captions=counters.table_captions,
            )
            words = render_markdown(Document(), path.read_text(encoding="utf-8"), file_counters, ROOT)
            problems.extend(f"{problem} ({path.name})" for problem in file_counters.problems)
            if path in body:
                print(f"  {path.name:40} {words:>6,} words")
        print(f"figures: {len(counters.figure_captions)}, tables: {len(counters.table_captions)}")
        if problems:
            print(f"PROBLEMS ({len(problems)}):")
            for problem in dict.fromkeys(problems):
                print("  -", problem)
            raise SystemExit(1)
        print("no problems found")
        return META_PATH

    document = Document()
    configure_styles(document)
    configure_page(document, meta["student_id"])
    document.core_properties.author = meta["student_id"]
    document.core_properties.last_modified_by = meta["student_id"]
    document.core_properties.title = meta["title"]
    document.core_properties.subject = f"{meta['module_code']} {meta['assessment']}"
    document.core_properties.comments = ""
    document.core_properties.keywords = ""

    add_cover(document, meta, body_words)
    add_toc(document, counters)
    add_abbreviations(document, meta, counters)
    for path in body + references + appendices:
        render_markdown(document, path.read_text(encoding="utf-8"), counters, ROOT)
    set_update_fields_on_open(document)

    output = output or REPORT_DIR / f"MLDA_CW1_{meta['student_id']}_report.docx"
    document.save(output)
    print(f"written {output} ({output.stat().st_size / 1e6:.1f} MB); body words {body_words:,}; "
          f"figures {len(counters.figure_captions)}; tables {len(counters.table_captions)}")
    problems = source_problems + sorted(set(counters.problems))
    if problems:
        print(f"PROBLEMS ({len(problems)}):")
        for problem in problems:
            print("  -", problem)
    return output


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--check", action="store_true", help="report word counts and cross-reference problems only")
    parser.add_argument("--output", type=Path, default=None)
    args = parser.parse_args()
    try:
        build(args.output, args.check)
    except SystemExit as error:
        if isinstance(error.code, str):
            print(error, file=sys.stderr)
        raise
